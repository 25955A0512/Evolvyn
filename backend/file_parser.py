"""Private, bounded document extraction. No uploaded content is executed."""
import io
import multiprocessing
import resource
import zipfile
from pathlib import Path

SUPPORTED={'.txt','.csv','.pdf','.docx'}

def _worker(raw,suffix,pipe):
    try:
        import os,logging
        os.environ.clear()
        logging.disable(logging.CRITICAL)
        resource.setrlimit(resource.RLIMIT_CPU,(4,5))
        resource.setrlimit(resource.RLIMIT_AS,(512*1024*1024,512*1024*1024))
        if suffix=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(raw),strict=True)
            if reader.is_encrypted:raise ValueError('Encrypted PDFs are unsupported')
            if len(reader.pages)>100:raise ValueError('Maximum 100 PDF pages')
            text='\n\n'.join(f'[Page {i+1}]\n'+(page.extract_text() or '') for i,page in enumerate(reader.pages))
        else:
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(f.file_size for f in archive.infolist())>10_000_000:raise ValueError('DOCX expands beyond 10 MB')
                if any(f.flag_bits&1 for f in archive.infolist()):raise ValueError('Encrypted DOCX is unsupported')
            document=Document(io.BytesIO(raw))
            text='\n'.join(p.text for p in document.paragraphs)
            text+='\n'+'\n'.join(' | '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        if not text.strip():raise ValueError('No extractable text; scanned PDFs need OCR')
        pipe.send({'text':text[:200_000]})
    except Exception:
        pipe.send({'error':'Document could not be safely parsed. Use an unencrypted text PDF, DOCX, or UTF-8 TXT/CSV.'})
    finally:pipe.close()

def extract(raw,name):
    suffix=Path(name).suffix.lower()
    if suffix not in SUPPORTED:raise ValueError('Supported formats: TXT, CSV, PDF and DOCX')
    if suffix in ('.txt','.csv'):
        try:return raw.decode('utf-8-sig')
        except UnicodeDecodeError:raise ValueError('TXT and CSV must use UTF-8 encoding')
    context=multiprocessing.get_context('spawn')
    parent,child=context.Pipe(duplex=False)
    process=context.Process(target=_worker,args=(raw,suffix,child))
    process.start();child.close()
    try:
        if not parent.poll(7):raise ValueError('Document extraction exceeded its processing budget')
        try:result=parent.recv()
        except EOFError:raise ValueError('Document extraction failed within its resource limit')
        if 'error' in result:raise ValueError(result['error'])
        return result['text']
    finally:
        parent.close()
        if process.is_alive():process.terminate()
        process.join(timeout=2)
